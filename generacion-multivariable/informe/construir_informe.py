"""Construye el informe DOCX a partir de la plantilla y de los resultados del experimento.

Uso (desde la raíz del proyecto):
    python generar_evidencias.py
    python informe/construir_informe.py [--capturas DIR] [--salida DIR]

Las capturas de la aplicación se leen de informe/capturas (excluida de git) y quedan
incrustadas en el DOCX, que se escribe en informe/. Requiere pandoc (https://pandoc.org) y python-docx.
"""

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import numpy as np
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

INFORME = Path(__file__).resolve().parent
RAIZ = INFORME.parent
sys.path.insert(0, str(RAIZ))

import config  # noqa: E402

FUENTE = "Times New Roman"


# --- Relleno de la plantilla -----------------------------------------------

def f(x, d=4):
    return f"{x:,.{d}f}"


def matriz_latex(m, d=4):
    filas = [" & ".join(f"{v:.{d}f}".rstrip("0").rstrip(".") if v != 0 else "0" for v in fila) for fila in m]
    return r"\begin{bmatrix}" + r"\\".join(filas) + r"\end{bmatrix}"


def tabla(encabezado, filas, alineacion):
    lineas = ["| " + " | ".join(encabezado) + " |",
              "|" + "|".join(":---" if a == "i" else "---:" for a in alineacion) + "|"]
    lineas += ["| " + " | ".join(fila) + " |" for fila in filas]
    return "\n".join(lineas)


def contar_pruebas():
    salida = subprocess.run([sys.executable, "-m", "pytest", "--collect-only", "-q", "-p", "no:cacheprovider"],
                            cwd=RAIZ, capture_output=True, text=True).stdout
    return re.search(r"(\d+) tests? collected", salida).group(1)


def valores(d):
    sigma = np.array(d["Sigma"])
    cov = np.array(d["cov_muestral"])
    r = np.array(d["R"])
    corr = np.array(d["corr_muestral"])
    mu = d["parametros"]["medias"]
    p = len(mu)

    momentos = [[f"$X_{i + 1}$", f(mu[i], 0), f(d["media_muestral"][i]), f(sigma[i, i], 0), f(cov[i, i])]
                for i in range(p)]
    covs = []
    for i in range(p):
        for j in range(i + 1, p):
            covs.append([f"$(X_{i + 1},X_{j + 1})$", f(sigma[i, j], 2), f(cov[i, j]), f(r[i, j], 2), f(corr[i, j])])
    conv = [[f"{c['n']:,}", f(c["error_medias"]), f(c["error_cov"])] for c in d["convergencia"]]

    ee = np.sqrt(d["tasa_teorica"] * (1 - d["tasa_teorica"]) / d["pares_propuestos"])
    vp = np.linalg.eigvalsh(sigma)

    return {
        "universidad": config.CURSO["universidad"],
        "facultad": config.CURSO["facultad"],
        "codigo": config.CURSO["codigo"],
        "curso": config.CURSO["nombre"],
        "seccion": config.CURSO["seccion"],
        "titulo": config.CURSO["titulo"],
        "fecha": config.CURSO["fecha"],
        "nombre": config.ESTUDIANTE["nombre"],
        "carne": config.ESTUDIANTE["carne"],
        "n": f"{d['parametros']['n']:,}",
        "np": f"{d['parametros']['n'] * p:,}",
        "semilla": str(d["parametros"]["semilla"]),
        "generador": d["generador_uniforme"],
        "L": matriz_latex(np.array(d["L"])),
        "valores_propios": ", ".join(f(v) for v in vp),
        "err_rec": f"{d['error_reconstruccion']:.1e}" if d["error_reconstruccion"] else "0 (exacto en punto flotante)",
        "aceptados": f"{d['pares_aceptados']:,}",
        "propuestos": f"{d['pares_propuestos']:,}",
        "tasa": f(d["tasa_aceptacion"]),
        "dif_tasa": f"{d['tasa_aceptacion'] - d['tasa_teorica']:+.4f}",
        "ee_tasa": f(ee),
        "media_z": f(d["media_z"]),
        "var_z": f(d["varianza_z"]),
        "err_medias": f(d["error_max_medias"]),
        "err_cov": f(d["error_max_cov"]),
        "err_corr": f(d["error_max_corr"]),
        "num_pruebas": contar_pruebas(),
        "tabla_momentos": tabla(["Variable", "Media teórica", "Media muestral", "Varianza teórica", "Varianza muestral"],
                                momentos, "iirrr"),
        "tabla_cov": tabla(["Pareja", "Covarianza teórica", "Covarianza muestral", "Correlación teórica",
                            "Correlación muestral"], covs, "irrrr"),
        "tabla_convergencia": tabla(["n", "Error máximo de medias", "Error máximo de covarianzas"], conv, "rrr"),
    }


def rellenar(plantilla, vals):
    texto = re.sub(r"\{\{(\w+)\}\}", lambda m: vals[m.group(1)], plantilla)
    pendientes = re.findall(r"\{\{\w+\}\}", texto)
    if pendientes:
        raise SystemExit(f"Marcadores sin valor: {pendientes}")
    return texto


# --- Formato del DOCX --------------------------------------------------------

def fuente(estilo, tam=None, negrita=None, cursiva=None):
    estilo.font.name = FUENTE
    estilo.font.color.rgb = RGBColor(0, 0, 0)
    rpr = estilo.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for atributo in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(atributo), FUENTE)
    for tema in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        rfonts.attrib.pop(qn(tema), None)
    if tam:
        estilo.font.size = Pt(tam)
    if negrita is not None:
        estilo.font.bold = negrita
    if cursiva is not None:
        estilo.font.italic = cursiva


def estilos(doc):
    s = doc.styles
    for nombre in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if nombre in [x.name for x in s]:
            fuente(s[nombre], 12)
            pf = s[nombre].paragraph_format
            pf.line_spacing = 1.5
            pf.space_after = Pt(6)
            pf.space_before = Pt(0)
            if nombre in ("Body Text", "First Paragraph"):
                pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for nivel, tam in ((1, 14), (2, 12)):
        h = s[f"Heading {nivel}"]
        fuente(h, tam, negrita=True, cursiva=False)
        h.paragraph_format.space_before = Pt(18 if nivel == 1 else 12)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
    for nombre in ("Image Caption", "Table Caption"):
        if nombre in [x.name for x in s]:
            fuente(s[nombre], 10, cursiva=False)
            s[nombre].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if nombre == "Image Caption" else WD_ALIGN_PARAGRAPH.LEFT
            s[nombre].paragraph_format.space_after = Pt(12)
            s[nombre].paragraph_format.line_spacing = 1.0
    s["Table Caption"].paragraph_format.keep_with_next = True
    if "Captioned Figure" in [x.name for x in s]:
        s["Captioned Figure"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        s["Captioned Figure"].paragraph_format.keep_with_next = True
    if "Source Code" in [x.name for x in s]:
        s["Source Code"].paragraph_format.line_spacing = 1.0
        s["Source Code"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if "Verbatim Char" in [x.name for x in s]:
        s["Verbatim Char"].font.size = Pt(9.5)
        s["Verbatim Char"].font.name = "Courier New"

    portada = s["Portada"]
    fuente(portada, 13)
    portada.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    portada.paragraph_format.line_spacing = 1.5
    titulo = s["Portada titulo"]
    fuente(titulo, 18, negrita=True)
    titulo.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo.paragraph_format.space_before = Pt(110)
    titulo.paragraph_format.space_after = Pt(110)
    titulo.paragraph_format.line_spacing = 1.3

    ref = s["Referencia"]
    fuente(ref, 12)
    ref.paragraph_format.left_indent = Inches(0.5)
    ref.paragraph_format.first_line_indent = Inches(-0.5)
    ref.paragraph_format.line_spacing = 1.5
    ref.paragraph_format.space_after = Pt(6)
    ref.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT


def campo_pagina(parrafo):
    run = parrafo.add_run()
    for tipo, texto in (("begin", None), (None, "PAGE"), ("separate", None), (None, "1"), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
            run._r.append(el)
        elif texto == "PAGE":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = " PAGE "
            run._r.append(el)
        else:
            el = OxmlElement("w:t")
            el.text = texto
            run._r.append(el)
    run.font.name = FUENTE
    run.font.size = Pt(11)


def pagina(doc):
    for sec in doc.sections:
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, lado, Inches(1))
        sec.different_first_page_header_footer = True
        pie = sec.footer.paragraphs[0]
        pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
        campo_pagina(pie)


def tablas(doc):
    for t in doc.tables:
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblpr = t._tbl.tblPr
        bordes = OxmlElement("w:tblBorders")
        for lado, grosor in (("top", 8), ("bottom", 8), ("insideH", 2)):
            b = OxmlElement(f"w:{lado}")
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), str(grosor))
            b.set(qn("w:color"), "000000")
            bordes.append(b)
        tblpr.append(bordes)
        for fila_idx, fila in enumerate(t.rows):
            for celda in fila.cells:
                for p in celda.paragraphs:
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.space_before = Pt(2)
                    for r in p.runs:
                        r.font.size = Pt(10.5)
                        r.font.name = FUENTE
                        if fila_idx == 0:
                            r.font.bold = True
        siguiente = t._tbl.getnext()
        if siguiente is not None and siguiente.tag == qn("w:p"):
            ppr = siguiente.get_or_add_pPr()
            espacio = OxmlElement("w:spacing")
            espacio.set(qn("w:before"), "240")
            ppr.append(espacio)
        # Celdas de encabezado: borde inferior marcado
        for celda in t.rows[0].cells:
            tcpr = celda._tc.get_or_add_tcPr()
            tcb = OxmlElement("w:tcBorders")
            b = OxmlElement("w:bottom")
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), "6")
            b.set(qn("w:color"), "000000")
            tcb.append(b)
            tcpr.append(tcb)


def leyendas(doc):
    """Pone en negrita el prefijo 'Figura n.' o 'Tabla n.' de cada leyenda."""
    for p in doc.paragraphs:
        if p.style.name not in ("Image Caption", "Table Caption") or not p.runs:
            continue
        primero = p.runs[0]
        m = re.match(r"((?:Figura|Tabla) \d+\.)(.*)", primero.text, re.S)
        if m:
            primero.text = m.group(2)
            nuevo = p.runs[0]._r
            negrita = OxmlElement("w:r")
            rpr = OxmlElement("w:rPr")
            rpr.append(OxmlElement("w:b"))
            negrita.append(rpr)
            t = OxmlElement("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = m.group(1)
            negrita.append(t)
            nuevo.addprevious(negrita)


def propiedades(doc):
    cp = doc.core_properties
    cp.title = config.CURSO["titulo"]
    cp.author = config.ESTUDIANTE["nombre"]
    cp.subject = f"{config.CURSO['codigo']} {config.CURSO['nombre']}, sección {config.CURSO['seccion']}"
    cp.keywords = "método polar, Marsaglia, normal multivariable, Cholesky, simulación"
    cp.language = "es-GT"
    idioma = doc.styles.element.find(qn("w:docDefaults"))
    if idioma is not None:
        rpr = idioma.find(qn("w:rPrDefault") + "/" + qn("w:rPr"))
        if rpr is not None:
            lang = rpr.find(qn("w:lang"))
            if lang is None:
                lang = OxmlElement("w:lang")
                rpr.append(lang)
            lang.set(qn("w:val"), "es-GT")


CAPTURAS = ("app_resultados.png", "app_graficas.png", "app_datos.png", "app_pizarra.png", "app_error.png")


def argumentos():
    a = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    a.add_argument("--capturas", type=Path, default=INFORME / "capturas",
                   help="Carpeta con las capturas de la aplicación (excluida de git).")
    a.add_argument("--salida", type=Path, default=INFORME,
                   help="Carpeta donde se escriben el DOCX y el PDF.")
    return a.parse_args()


def main():
    args = argumentos()
    if not (RAIZ / "resultados" / "resumen_referencia.json").exists():
        raise SystemExit("Primero ejecute: python generar_evidencias.py")
    faltan = [c for c in CAPTURAS if not (args.capturas / c).exists()]
    if faltan:
        raise SystemExit(f"Faltan capturas de la aplicación en {args.capturas}: {', '.join(faltan)}")
    args.salida.mkdir(parents=True, exist_ok=True)

    datos = json.loads((RAIZ / "resultados" / "resumen_referencia.json").read_text())
    md = rellenar((INFORME / "plantilla_informe.md").read_text(), valores(datos))

    with tempfile.TemporaryDirectory() as tmp:
        fuente_md = Path(tmp) / "informe.md"
        fuente_md.write_text(md)
        crudo = Path(tmp) / "informe_pandoc.docx"
        rutas = os.pathsep.join([str(INFORME), str(args.capturas.resolve())])
        subprocess.run(["pandoc", str(fuente_md), "-f", "markdown", "-o", str(crudo),
                        "--resource-path", rutas, "-M", "lang=es-GT"], cwd=INFORME, check=True)
        doc = Document(crudo)
    salida = args.salida / "Informe_Generacion_Multivariable.docx"
    construir(doc, salida)
    exportar_pdf(salida)


def construir(doc, salida):

    estilos(doc)
    pagina(doc)
    tablas(doc)
    leyendas(doc)
    propiedades(doc)
    doc.save(salida)
    print(f"Informe generado: {salida}")


def exportar_pdf(salida):
    soffice = shutil.which("soffice")
    if soffice:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(salida.parent), str(salida)],
                       check=True, capture_output=True)
        print(f"PDF generado: {salida.with_suffix('.pdf')}")
    else:
        print("LibreOffice no está instalado; exporte el PDF desde Word (Archivo > Guardar como > PDF).")


if __name__ == "__main__":
    main()
